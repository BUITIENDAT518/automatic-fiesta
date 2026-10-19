"""Order exports and the daily digest, shared by the API (CSV endpoint) and the Telegram bot (daily push)."""

from __future__ import annotations

import csv
import io
import os
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

CSV_COLUMNS = [
    ("created_at", lambda o: o["created_at"]), ("updated_at", lambda o: o["updated_at"]),
    ("platform", lambda o: o["platform"]), ("shop", lambda o: o.get("shop_name") or ""),
    ("order_id", lambda o: o["order_id"]), ("status", lambda o: o["status"]),
    ("raw_status", lambda o: o.get("raw_status") or ""),
    ("total_amount", lambda o: o["total_amount"]), ("currency", lambda o: o["currency"]),
    ("customer", lambda o: (o["customer"].get("name") or "")), ("phone", lambda o: (o["customer"].get("phone") or "")),
    ("recipient", lambda o: (o["shipping"].get("recipient_name") or "")),
    ("address", lambda o: (o["shipping"].get("address") or "")),
    ("province", lambda o: (o["shipping"].get("province") or "")),
    ("region", lambda o: (o["shipping"].get("region") or "")),
    ("carrier", lambda o: (o["shipping"].get("carrier") or "")),
    ("tracking_number", lambda o: (o["shipping"].get("tracking_number") or "")),
    ("logistics_status", lambda o: (o["shipping"].get("logistics_status") or "")),
    ("items", lambda o: "; ".join(f"{i.get('name')} x{i.get('quantity')}" for i in o.get("items") or [])),
]


def rows_to_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    buf.write("\ufeff")  # BOM so Excel opens UTF-8 Vietnamese correctly
    w = csv.writer(buf)
    w.writerow([name for name, _ in CSV_COLUMNS])
    for o in rows:
        w.writerow([fn(o) for _, fn in CSV_COLUMNS])
    return buf.getvalue()


def rows_to_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append([name for name, _ in CSV_COLUMNS])
    for o in rows:
        ws.append([fn(o) for _, fn in CSV_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def append_xlsx(path: Path, rows: list[dict]) -> int:
    """Append rows to the master XLSX (creating it with the header on first run); returns rows added.

    The file is rewritten to a sibling temp file then atomically moved over the original.
    """
    if not rows and path.exists():
        return 0
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.worksheets[0]
    else:
        wb = Workbook()
        ws = wb.active
        ws.append([name for name, _ in CSV_COLUMNS])
    added = 0
    for o in rows:
        ws.append([fn(o) for _, fn in CSV_COLUMNS])
        added += 1
    tmp = path.with_name(path.name + ".tmp")
    wb.save(tmp)
    os.replace(tmp, path)
    return added


def digest_text(stats: dict, hours: int) -> str:
    def top(dim: str, n: int = 5) -> str:
        items = stats.get(f"by_{dim}") or []
        return "\n".join(f"  - {r['key']}: {r['orders']:,} đơn" for r in items[:n]) or "  (chưa có)"

    revenue = f"{stats.get('revenue', 0):,.0f} đ"
    delivered = next((r["orders"] for r in stats.get("by_status") or [] if r["key"] == "delivered"), 0)
    return (
        f"Báo cáo đơn hàng {hours} giờ gần nhất ({datetime.now(timezone.utc).astimezone().strftime('%d/%m/%Y %H:%M')})\n"
        f"Tổng đơn: {stats.get('orders', 0):,} · Doanh thu: {revenue} · Đã giao: {delivered:,}\n\n"
        f"Theo sàn:\n{top('platform')}\n\n"
        f"Theo vùng miền:\n{top('region', 3)}\n\n"
        f"Theo trạng thái:\n{top('status', 7)}"
    )


async def build_daily(db, settings, hours: int, reveal_pii: bool) -> tuple[str, bytes]:
    stats = await db.stats(hours, {"platform": None, "shop_id": None, "region": None,
                                   "province": None, "status": None, "carrier": None})
    _, rows = await db.query_orders(hours, {}, None, limit=100_000, offset=0, reveal_pii=reveal_pii)
    return digest_text(stats, hours), rows_to_csv(rows).encode("utf-8")
