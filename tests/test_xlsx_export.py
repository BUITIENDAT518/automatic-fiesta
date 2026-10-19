import io

from openpyxl import load_workbook

from orderhub.report import CSV_COLUMNS, append_xlsx, rows_to_xlsx

HEADERS = [name for name, _ in CSV_COLUMNS]
ORDER_ID_COL = HEADERS.index("order_id") + 1
ITEMS_COL = HEADERS.index("items") + 1


def _row(order_id, platform="shopee", shop_id="100001", created="2026-10-19T01:00:00+00:00"):
    return {
        "created_at": created, "updated_at": created, "platform": platform, "shop_id": shop_id,
        "shop_name": "Shop A", "order_id": order_id, "status": "pending", "raw_status": "UNPAID",
        "total_amount": 150000.0, "currency": "VND",
        "customer": {"id": "c1", "name": "Nguyễn Văn A", "phone": "0901234567", "email": "a@x.vn"},
        "shipping": {"recipient_name": "Nguyễn Văn A", "phone": "0901234567",
                     "address": "1 Lê Lợi, P. Bến Nghé, Q1", "ward": "Bến Nghé", "district": "Quận 1",
                     "province": "Hồ Chí Minh", "province_legacy": None, "region": "Nam",
                     "subregion": "Đông Nam Bộ", "carrier": "ghtk", "tracking_number": "TN001",
                     "logistics_status": "ready"},
        "items": [{"name": "Áo thun", "quantity": 2}, {"name": "Quần jean", "quantity": 1}],
    }


def _values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_rows_to_xlsx_header_and_cells():
    data = rows_to_xlsx([_row("O1"), _row("O2")])
    ws = load_workbook(io.BytesIO(data)).active
    vals = _values(ws)
    assert vals[0] == HEADERS
    assert len(vals) == 3
    assert vals[1][ORDER_ID_COL - 1] == "O1"
    assert vals[1][ITEMS_COL - 1] == "Áo thun x2; Quần jean x1"


def test_append_xlsx_creates_then_appends(tmp_path):
    path = tmp_path / "orders-master.xlsx"
    assert append_xlsx(path, [_row("O1")]) == 1
    assert append_xlsx(path, [_row("O2"), _row("O3", platform="tiktok", shop_id="7494")]) == 2
    ws = load_workbook(path).active
    vals = _values(ws)
    assert vals[0] == HEADERS
    assert [r[ORDER_ID_COL - 1] for r in vals[1:]] == ["O1", "O2", "O3"]
    assert len({(r[2], r[ORDER_ID_COL - 1]) for r in vals[1:]}) == 3  # không lặp đơn


def test_append_xlsx_empty_rows(tmp_path):
    path = tmp_path / "sub" / "m.xlsx"
    assert append_xlsx(path, []) == 0  # tạo file chỉ có header
    ws = load_workbook(path).active
    assert _values(ws) == [HEADERS]
    assert append_xlsx(path, []) == 0  # file đã có: không ghi lại
    assert append_xlsx(path, [_row("O1")]) == 1
    assert load_workbook(path).active.max_row == 2
